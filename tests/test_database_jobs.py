"""Durable execution identity, admission, revocation, complete paging and export."""

from base64 import b64encode
from contextlib import contextmanager
from copy import deepcopy
import hashlib
import io
import json
import threading
import time
from unittest.mock import Mock

from fastapi import HTTPException, Request
from fastapi.testclient import TestClient
from itsdangerous import TimestampSigner
from openpyxl import load_workbook
import pytest

from eng_platform_api.config import config
from eng_platform_api.main import app
from eng_platform_api.services import (
    database_console as console,
    database_job_store as store,
    database_jobs as jobs,
    database_result_store as results,
    database_sessions as sessions,
    database_tasks as tasks,
)
from eng_platform_api.services.database_registry import DatabaseUnavailable

HEADERS = {"Origin": "http://localhost:5173", "X-Requested-With": "EngineeringPlatform"}
DATABASE = {
    "id": "sample",
    "name": "Sample",
    "schemas": ["sample"],
    "allowed_logins": ["reader", "other"],
    "dsn_env": "ENG_PLATFORM_DATABASE_DSN_SAMPLE",
}
COLUMNS = [
    {"key": "c0", "name": "same", "data_type": "integer"},
    {"key": "c1", "name": "same", "data_type": "numeric"},
]


class Control:
    def __init__(self):
        self.records = {}
        self.lock = threading.RLock()

    def get(self, kind, identity):
        with self.lock:
            return deepcopy(self.records.get(store.key(kind, identity)))

    def mutate(self, references, transform):
        with self.lock:
            values = {
                store.key(*reference): deepcopy(self.records.get(store.key(*reference)))
                for reference in references
            }
            changes = transform(values)
            for label, value in changes.items():
                if value is None:
                    self.records.pop(label, None)
                else:
                    self.records[label] = deepcopy(value)

    def find(self, kind, field, value, *, limit):
        with self.lock:
            return deepcopy(
                [
                    record
                    for label, record in self.records.items()
                    if label.startswith(f"{kind}:")
                    and (
                        record.get(field, float("inf")) <= value
                        if field == "expires_at"
                        else record.get(field) == value
                    )
                ][:limit]
            )


class Objects:
    def __init__(self):
        self.values = {}
        self.generation = 0

    def put(self, path, data, content_type):
        if path in self.values:
            raise ValueError("Create-only object exists")
        self.generation += 1
        self.values[path] = (self.generation, bytes(data))
        return self.generation

    def read(self, reference, maximum):
        generation, data = self.values[reference["object"]]
        assert generation == reference["generation"]
        return data[: maximum + 1]

    def delete(self, reference):
        self.values.pop(reference["object"], None)

    def purge(self, identity):
        self.values = {
            name: value
            for name, value in self.values.items()
            if not name.startswith((f"inputs/{identity}/", f"results/{identity}/"))
        }

    def contains(self, identity):
        return any(
            path.startswith((f"inputs/{identity}/", f"results/{identity}/"))
            for path in self.values
        )

    @contextmanager
    def export_sink(self, path):
        sink = io.BytesIO()
        yield sink
        self.put(path, sink.getvalue(), "xlsx")

    def export_reference(self, path):
        generation, data = self.values[path]
        return {"object": path, "generation": generation, "bytes": len(data)}

    def download(self, reference, authorize):
        generation, data = self.values[reference["object"]]
        assert generation == reference["generation"]
        for index in range(0, len(data), results.CHUNK_BYTES):
            authorize()
            yield data[index : index + results.CHUNK_BYTES]


@pytest.fixture
def configured(monkeypatch, tmp_path):
    control, objects, deliveries = Control(), Objects(), []
    monkeypatch.setattr(config, "mock_mode", False)
    for field, value in {
        "enabled": True,
        "executions_enabled": True,
        "allowed_logins": ("reader", "other"),
        "project_id": "test-project",
        "collection": "eng_platform_database_control",
        "result_bucket": "private-test-bucket",
        "query_queue": "query",
        "export_queue": "export",
        "cleanup_queue": "cleanup",
        "worker_service_account": "tasks@example.iam.gserviceaccount.com",
        "api_origin": "https://api.example",
        "policy_version": "policy-1",
    }.items():
        monkeypatch.setattr(config.databases, field, value)
    monkeypatch.setattr(
        config.auth, "session_secret", "strong-test-session-secret-32-chars"
    )
    monkeypatch.setattr(config.auth, "github_client_id", "test-client")
    monkeypatch.setattr(config.auth, "github_client_secret", "test-secret")
    monkeypatch.setattr(config.auth, "frontend_url", HEADERS["Origin"])
    registry = tmp_path / "registry.json"
    registry.write_text(json.dumps({"databases": [DATABASE]}))
    monkeypatch.setattr(config.databases, "registry_path", str(registry))
    monkeypatch.setattr(
        config.databases,
        "registry_sha256",
        hashlib.sha256(registry.read_bytes()).hexdigest(),
    )
    monkeypatch.setenv(
        DATABASE["dsn_env"],
        "postgresql://reader:synthetic@127.0.0.1:55439/readonly_test",
    )
    monkeypatch.setattr(
        tasks,
        "_test_dispatcher",
        lambda kind, identity, schedule: deliveries.append((kind, identity, schedule)),
    )
    with store.testing_backend(control), results.testing_backend(objects):
        control.records["control:active_policy"] = {
            "kind": "control",
            "enabled": True,
            "version": "policy-1",
            "fingerprint": jobs.policy_fingerprint(),
        }
        nonce = sessions.create("reader")
        request = Request(
            {
                "type": "http",
                "headers": [],
                "session": {
                    "github_login": "reader",
                    "github_auth_provider": "github_oauth",
                    sessions.COOKIE_FIELD: nonce,
                },
            }
        )
        workspace = jobs.create_workspace(request, "sample")
        yield {
            "control": control,
            "objects": objects,
            "deliveries": deliveries,
            "request": request,
            "workspace": workspace,
            "nonce": nonce,
        }


def submit(environment, statement="SELECT 1", key="request-1"):
    return jobs.create_execution(
        environment["request"],
        "sample",
        environment["workspace"]["workspace_id"],
        statement,
        key,
    )


def stream_stub(monkeypatch, environment, count=1205, side_effect=None):
    def run(
        database,
        statement,
        authorize,
        on_columns,
        on_rows,
        *,
        on_connection,
        admission,
        timeout_seconds,
    ):
        assert not any(
            name.startswith("inputs/") for name in environment["objects"].values
        )
        assert 0 < timeout_seconds <= 240
        with admission:
            authorize()
            on_columns(COLUMNS)
            for start in range(0, count, 16):
                on_rows(
                    [
                        [index, "1234567890.1234567890123456789"]
                        for index in range(start, min(count, start + 16))
                    ]
                )
            if side_effect:
                side_effect()
        return {"row_count": count, "elapsed_ms": 10}

    mock = Mock(side_effect=run)
    monkeypatch.setattr(console, "stream_query", mock)
    return mock


def test_execution_is_once_and_every_fixed_page_and_export_use_identical_snapshot(
    configured, monkeypatch
):
    stream = stream_stub(monkeypatch, configured)
    value = submit(configured)
    identity = value["execution_id"]
    jobs.run_query(identity)
    jobs.run_query(identity)
    stream.assert_called_once()
    request, workspace_id = (
        configured["request"],
        configured["workspace"]["workspace_id"],
    )
    state = jobs.execution(request, "sample", workspace_id, identity)
    assert state["status"] == "completed" and state["row_count"] == 1205
    assert 3595 < state["expires_at"] - time.time() <= 3600
    for size in (25, 50, 100):
        page = jobs.page(request, "sample", workspace_id, identity, 3, size)
        assert page["rows"][0][0] == 3 * size
        assert len(page["rows"]) == size
        assert page["total_pages"] == (1205 + size - 1) // size
    last = jobs.page(request, "sample", workspace_id, identity, 12, 100)
    assert [row[0] for row in last["rows"]] == list(range(1200, 1205))
    exported = jobs.create_export(request, "sample", workspace_id, identity)
    assert (
        jobs.create_export(request, "sample", workspace_id, identity)["export_id"]
        == exported["export_id"]
    )
    jobs.run_export(exported["export_id"])
    jobs.run_export(exported["export_id"])
    response = jobs.export(
        request, "sample", workspace_id, identity, exported["export_id"]
    )
    assert response["status"] == "completed"
    output, size = jobs.open_download(
        request, "sample", workspace_id, identity, exported["export_id"]
    )
    data = b"".join(output)
    assert len(data) == size
    workbook = load_workbook(io.BytesIO(data), read_only=True)
    sheet = (
        workbook["Datos_1"]
        if "Datos_1" in workbook.sheetnames
        else workbook.worksheets[0]
    )
    cells = list(sheet.values)
    assert cells[0] == ("same", "same")
    assert len(cells) == 1206 and cells[-1] == (
        "1204",
        "1234567890.1234567890123456789",
    )
    assert stream.call_count == 1
    assert not configured["control"].records["control:budget"]["permits"]


def test_idempotency_does_not_cache_different_sql_or_reexecute_completed_result(
    configured, monkeypatch
):
    value = submit(configured)
    assert submit(configured)["execution_id"] == value["execution_id"]
    assert {item[1] for item in configured["deliveries"] if item[0] == "query"} == {
        value["execution_id"]
    }
    with pytest.raises(HTTPException) as error:
        submit(configured, "SELECT 2")
    assert error.value.status_code == 409
    assert len(configured["objects"].values) == 1


def test_previous_completed_snapshot_survives_another_failed_execution(
    configured, monkeypatch
):
    stream_stub(monkeypatch, configured, count=2)
    first = submit(configured)
    jobs.run_query(first["execution_id"])
    second = submit(configured, "SELECT 2", "request-2")
    monkeypatch.setattr(
        console, "stream_query", Mock(side_effect=DatabaseUnavailable("SECRET SQL dsn"))
    )
    jobs.run_query(second["execution_id"])
    assert (
        jobs.execution(
            configured["request"],
            "sample",
            first["workspace_id"],
            second["execution_id"],
        )["error_code"]
        == "QUERY_FAILED"
    )
    page = jobs.page(
        configured["request"],
        "sample",
        first["workspace_id"],
        first["execution_id"],
        0,
        100,
    )
    assert page["row_count"] == 2


@pytest.mark.parametrize("revocation", ["session", "workspace", "policy"])
def test_atomic_final_publish_rejects_revocation_after_last_worker_check(
    configured, monkeypatch, revocation
):
    stream_stub(monkeypatch, configured, count=2)
    value = submit(configured)
    original = jobs._finish

    def revoke_then_finish(identity, claim, state, **kwargs):
        if state == "completed":
            if revocation == "session":
                sessions.revoke(configured["request"])
            elif revocation == "workspace":
                configured["control"].records[f"workspace:{value['workspace_id']}"][
                    "state"
                ] = "purged"
            else:
                configured["control"].records["control:active_policy"]["enabled"] = (
                    False
                )
        return original(identity, claim, state, **kwargs)

    monkeypatch.setattr(jobs, "_finish", revoke_then_finish)
    jobs.run_query(value["execution_id"])
    state = store.get("execution", value["execution_id"])
    assert state["state"] == "cancelled" and state["manifest"] is None
    assert not configured["objects"].values


def test_cancelled_queued_execution_never_dispatches_sql_and_cleanup_releases_reservation(
    configured, monkeypatch
):
    stream = stream_stub(monkeypatch, configured)
    value = submit(configured)
    jobs.cancel(
        configured["request"], "sample", value["workspace_id"], value["execution_id"]
    )
    jobs.run_query(value["execution_id"])
    stream.assert_not_called()
    jobs.cleanup(value["execution_id"])
    assert not configured["objects"].values
    assert not configured["control"].records["control:budget"]["reservations"]


def test_purge_workspace_invalidates_old_results_and_removes_all_objects(
    configured, monkeypatch
):
    stream_stub(monkeypatch, configured, count=2)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    jobs.purge_workspace(configured["request"], "sample", value["workspace_id"])
    with pytest.raises(HTTPException) as error:
        jobs.page(
            configured["request"],
            "sample",
            value["workspace_id"],
            value["execution_id"],
            0,
            100,
        )
    assert error.value.status_code == 410
    jobs.cleanup(value["workspace_id"])
    assert not configured["objects"].values
    assert not configured["control"].records["control:budget"]["reservations"]


def test_lost_worker_never_reclaims_execution_or_runs_sql(configured, monkeypatch):
    stream = stream_stub(monkeypatch, configured)
    value = submit(configured)
    claimed = jobs._claim("execution", value["execution_id"])
    configured["control"].records[f"execution:{value['execution_id']}"][
        "worker_deadline"
    ] = 0
    jobs.sweep()
    jobs.run_query(value["execution_id"])
    stream.assert_not_called()
    state = store.get("execution", value["execution_id"])
    assert state["state"] == "failed" and state["error_code"] == "WORKER_LOST"
    assert claimed["claim"]
    assert not configured["objects"].values


def test_roles_allow_one_query_and_one_metadata_connection_and_reads_are_separate(
    configured,
):
    database = jobs._database("sample", "reader")
    with jobs.connection_slot(database, "query", "reader"):
        with jobs.connection_slot(database, "metadata"):
            with pytest.raises(HTTPException):
                with jobs.connection_slot(database, "metadata"):
                    pass
        with pytest.raises(HTTPException) as denied:
            with jobs.connection_slot(database, "query", "other"):
                pass
        assert denied.value.status_code == 429
        with jobs._permit("read", "reader"):
            with jobs._permit("read", "other"):
                with pytest.raises(HTTPException):
                    with jobs._permit("read", "third"):
                        pass
    assert not configured["control"].records["control:budget"]["permits"]


def test_storage_budget_reserves_before_dispatch_and_remains_until_cleanup_ack(
    configured, monkeypatch
):
    value = submit(configured)
    second = submit(configured, "SELECT 2", "second")
    with pytest.raises(HTTPException) as denied:
        submit(configured, "SELECT 3", "third")
    assert denied.value.status_code == 429
    monkeypatch.setattr(
        results, "purge", Mock(side_effect=DatabaseUnavailable("storage outage"))
    )
    jobs.cancel(
        configured["request"], "sample", value["workspace_id"], value["execution_id"]
    )
    with pytest.raises(DatabaseUnavailable):
        jobs.cleanup(value["execution_id"])
    assert (
        value["execution_id"]
        in configured["control"].records["control:budget"]["reservations"]
    )
    assert second["status"] == "queued"


def test_workspaces_outlive_one_hour_but_results_expire_one_hour_after_completion(
    configured, monkeypatch
):
    assert configured["workspace"]["expires_at"] > time.time() + 11 * 3600
    stream_stub(monkeypatch, configured, count=1)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    configured["control"].records[f"execution:{value['execution_id']}"][
        "expires_at"
    ] = time.time() - 1
    jobs.cleanup(value["execution_id"])
    with pytest.raises(HTTPException) as error:
        jobs.page(
            configured["request"],
            "sample",
            value["workspace_id"],
            value["execution_id"],
            0,
            100,
        )
    assert error.value.status_code == 410
    assert (
        jobs.workspace(configured["request"], "sample", value["workspace_id"])[
            "workspace_id"
        ]
        == value["workspace_id"]
    )


def browser(environment):
    client = TestClient(app)
    secret = next(
        item for item in app.user_middleware if item.cls.__name__ == "SessionMiddleware"
    ).kwargs["secret_key"]
    client.cookies.set(
        "session",
        TimestampSigner(secret)
        .sign(b64encode(json.dumps(environment["request"].session).encode()))
        .decode(),
    )
    return client


def test_http_contract_private_body_fixed_page_validation_no_store_and_old_cookie_denied(
    configured, monkeypatch
):
    client = browser(configured)
    response = client.get("/api/databases")
    assert response.status_code == 200 and response.json()["workspace_enabled"] is True
    assert response.json()["databases"][0]["max_rows"] == 500
    assert response.json()["databases"][0]["timeout_seconds"] == 5
    assert response.json()["workspace_limits"]["timeout_seconds"] == 240
    route = f"/api/databases/sample/workspaces/{configured['workspace']['workspace_id']}/executions"
    bad = client.post(
        route,
        json={"sql": "SELECT SECRET", "client_request_id": "valid", "extra": "SECRET"},
        headers=HEADERS,
    )
    assert (
        bad.status_code == 422
        and "SECRET" not in bad.text
        and bad.headers["cache-control"] == "no-store"
    )
    created = client.post(
        route, json={"sql": "SELECT 1", "client_request_id": "valid"}, headers=HEADERS
    )
    assert created.status_code == 202
    stream_stub(monkeypatch, configured, count=1)
    jobs.run_query(created.json()["execution_id"])
    response = client.post(
        f"{route}/{created.json()['execution_id']}/pages",
        json={"page_index": 0, "page_size": 200},
        headers=HEADERS,
    )
    assert response.status_code == 422
    old = Request(
        {
            "type": "http",
            "headers": [],
            "session": {
                "github_login": "reader",
                "github_auth_provider": "github_oauth",
            },
        }
    )
    configured["request"] = old
    assert browser(configured).get("/api/databases").status_code == 401


def test_native_download_is_authenticated_without_ajax_header_and_blocks_cross_site(
    configured, monkeypatch
):
    stream_stub(monkeypatch, configured, count=1)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    exported = jobs.create_export(
        configured["request"], "sample", value["workspace_id"], value["execution_id"]
    )
    jobs.run_export(exported["export_id"])
    path = jobs.export(
        configured["request"],
        "sample",
        value["workspace_id"],
        value["execution_id"],
        exported["export_id"],
    )["download_path"]
    client = browser(configured)
    response = client.get(path, headers={"Sec-Fetch-Site": "same-origin"})
    assert response.status_code == 200 and response.content.startswith(b"PK")
    assert (
        response.headers["cache-control"] == "no-store"
        and "attachment" in response.headers["content-disposition"]
    )
    assert client.get(path, headers={"Sec-Fetch-Site": "cross-site"}).status_code == 403
    sessions.revoke(configured["request"])
    assert client.get(path).status_code == 401


def test_running_duplicate_delivery_cannot_reclaim_the_cursor(configured):
    value = submit(configured)
    claims = []

    def claim():
        claims.append(jobs._claim("execution", value["execution_id"]))

    threads = [threading.Thread(target=claim) for _ in range(8)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert sum(result is not None for result in claims) == 1
    assert jobs._claim("execution", value["execution_id"]) is None


def test_watcher_cancels_blocked_connection_after_logout(configured):
    value = submit(configured)
    claimed = jobs._claim("execution", value["execution_id"])
    context = jobs.RunContext(claimed)
    cancelled = threading.Event()
    connection = Mock()
    connection.cancel_safe.side_effect = lambda **kwargs: cancelled.set()
    context.on_connection(connection)
    context.authorize(force=True)
    with context.watching():
        sessions.revoke(configured["request"])
        assert cancelled.wait(3)
        with pytest.raises(HTTPException):
            context.authorize()
    connection.cancel_safe.assert_called_once_with(timeout=2)


def test_callback_authorization_uses_local_cache_but_forced_watcher_revalidates(
    configured, monkeypatch
):
    value = submit(configured)
    context = jobs.RunContext(jobs._claim("execution", value["execution_id"]))
    original = sessions.validate
    checks = Mock(side_effect=original)
    monkeypatch.setattr(sessions, "validate", checks)
    context.authorize(force=True)
    for _ in range(500):
        context.authorize()
    assert checks.call_count == 1
    context.authorize(force=True)
    assert checks.call_count == 2


def test_export_worker_checks_source_ttl_while_uploading(configured, monkeypatch):
    stream_stub(monkeypatch, configured, count=2)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    exported = jobs.create_export(
        configured["request"], "sample", value["workspace_id"], value["execution_id"]
    )
    context = jobs.RunContext(
        jobs._claim("export", exported["export_id"]), kind="export"
    )
    context.authorize(force=True)
    configured["control"].records[f"execution:{value['execution_id']}"][
        "expires_at"
    ] = time.time() - 1
    with pytest.raises(HTTPException):
        context.authorize(force=True)


def test_failed_export_can_retry_same_snapshot_after_physical_cleanup(
    configured, monkeypatch
):
    from eng_platform_api.services import database_xlsx

    stream = stream_stub(monkeypatch, configured, count=2)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    exported = jobs.create_export(
        configured["request"], "sample", value["workspace_id"], value["execution_id"]
    )
    original = database_xlsx.write_xlsx
    monkeypatch.setattr(
        database_xlsx, "write_xlsx", Mock(side_effect=HTTPException(422, "budget"))
    )
    jobs.run_export(exported["export_id"])
    failed = store.get("export", exported["export_id"])
    assert failed["state"] == "failed" and failed["artifact_cleaned"]
    assert (
        exported["export_id"]
        not in configured["control"].records["control:budget"]["reservations"]
    )
    monkeypatch.setattr(database_xlsx, "write_xlsx", original)
    retried = jobs.create_export(
        configured["request"], "sample", value["workspace_id"], value["execution_id"]
    )
    assert retried["export_id"] != exported["export_id"]
    jobs.run_export(retried["export_id"])
    assert store.get("export", retried["export_id"])["state"] == "completed"
    stream.assert_called_once()


def test_cleanup_dispatch_failure_cannot_destroy_published_result(
    configured, monkeypatch
):
    stream_stub(monkeypatch, configured, count=2)
    value = submit(configured)
    monkeypatch.setattr(
        tasks, "enqueue", Mock(side_effect=DatabaseUnavailable("tasks outage"))
    )
    with pytest.raises(DatabaseUnavailable):
        jobs.run_query(value["execution_id"])
    state = store.get("execution", value["execution_id"])
    assert state["state"] == "completed"
    assert results.manifest(state["manifest"])["row_count"] == 2
    jobs.run_query(value["execution_id"])


def test_staging_cleanup_is_scheduled_before_upload_and_dispatch_outage_stages_nothing(
    configured, monkeypatch
):
    staged = Mock(side_effect=results.stage_sql)
    enqueue = Mock(side_effect=DatabaseUnavailable("unavailable"))
    monkeypatch.setattr(results, "stage_sql", staged)
    monkeypatch.setattr(tasks, "enqueue", enqueue)
    with pytest.raises(DatabaseUnavailable):
        submit(configured)
    staged.assert_not_called()
    assert not configured["objects"].values


def test_cleanup_missing_control_document_erases_orphan_sql(configured):
    identity = "d" * 32
    results.stage_sql(identity, "SELECT 1")
    jobs.cleanup(identity)
    assert not configured["objects"].values


def test_running_purge_keeps_reservations_until_worker_deadline_and_retries(configured):
    value = submit(configured)
    jobs._claim("execution", value["execution_id"])
    jobs.purge_workspace(configured["request"], "sample", value["workspace_id"])
    with pytest.raises(DatabaseUnavailable):
        jobs.cleanup(value["workspace_id"])
    assert (
        value["execution_id"]
        in configured["control"].records["control:budget"]["reservations"]
    )
    configured["control"].records[f"execution:{value['execution_id']}"][
        "worker_deadline"
    ] = 0
    jobs.cleanup(value["workspace_id"])
    assert not configured["objects"].values
    assert not configured["control"].records["control:budget"]["reservations"]


def test_sweeper_makes_progress_past_cleaned_expired_metadata(configured):
    # Two scan windows: cleaned records move out of the current expired range.
    now = time.time()
    for index in range(70):
        identity = f"{index:032x}"
        configured["control"].records[f"execution:{identity}"] = {
            "kind": "execution",
            "execution_id": identity,
            "workspace_id": configured["workspace"]["workspace_id"],
            "state": "failed",
            "expires_at": now - 1,
            "worker_deadline": 0,
        }
    jobs.sweep()
    jobs.sweep()
    records = store.find("execution", "expires_at", now, limit=256)
    assert not records
    assert (
        sum(
            value.get("cleanup_done", False)
            for value in configured["control"].records.values()
        )
        == 70
    )


def test_query_admission_global_owner_and_role_limits_are_distributed(configured):
    with jobs._permit("query", "reader", role="role-a"):
        with pytest.raises(HTTPException):
            with jobs._permit("query", "reader", role="role-b"):
                pass
        with jobs._permit("query", "other", role="role-b"):
            with pytest.raises(HTTPException):
                with jobs._permit("query", "third", role="role-c"):
                    pass
    assert not configured["control"].records["control:budget"]["permits"]


def test_legacy_query_contract_remains_available_during_workspace_rollout(
    configured, monkeypatch
):
    query = Mock(
        return_value={
            "columns": ["value"],
            "rows": [["1"]],
            "row_count": 1,
            "truncated": False,
            "elapsed_ms": 1,
        }
    )
    monkeypatch.setattr(console, "query", query)
    response = browser(configured).post(
        "/api/databases/sample/query",
        json={"sql": "SELECT 1", "max_rows": 100},
        headers=HEADERS,
    )
    assert response.status_code == 200 and response.json()["truncated"] is False
    assert query.call_args.kwargs["admission"] is not None


def test_queued_outbox_repairs_crash_before_dispatch_without_repeating_running_sql(
    configured,
):
    value = submit(configured)
    configured["deliveries"].clear()
    repeated = submit(configured)
    assert repeated["execution_id"] == value["execution_id"]
    assert configured["deliveries"] == [("query", value["execution_id"], None)]
    configured["deliveries"].clear()
    jobs.sweep()
    assert configured["deliveries"] == [("query", value["execution_id"], None)]
    jobs._claim("execution", value["execution_id"])
    configured["deliveries"].clear()
    jobs.sweep()
    assert not configured["deliveries"]


def test_queued_query_outbox_stops_after_dispatch_deadline_without_querying(
    configured, monkeypatch
):
    stream = stream_stub(monkeypatch, configured)
    value = submit(configured)
    configured["control"].records[f"execution:{value['execution_id']}"][
        "dispatch_deadline"
    ] = time.time() - 1
    jobs.sweep()
    current = store.get("execution", value["execution_id"])
    assert current["error_code"] == "DISPATCH_FAILED" and current["state"] == "expired"
    assert not configured["objects"].values
    assert not configured["control"].records["control:budget"]["reservations"]
    jobs.run_query(value["execution_id"])
    stream.assert_not_called()


def test_dispatch_ack_uncertainty_never_erases_or_releases_running_capture(
    configured, monkeypatch
):
    enqueue = tasks.enqueue

    def ambiguous(kind, identity, *, schedule_at=None):
        if kind == "query":
            jobs._claim("execution", identity)
            raise DatabaseUnavailable("ACK uncertain")
        return enqueue(kind, identity, schedule_at=schedule_at)

    monkeypatch.setattr(tasks, "enqueue", ambiguous)
    with pytest.raises(DatabaseUnavailable):
        submit(configured)
    value = store.find("execution", "state", "running", limit=1)[0]
    assert (
        value["execution_id"]
        in configured["control"].records["control:budget"]["reservations"]
    )
    assert value["input"]["object"] in configured["objects"].values


def test_logout_marks_every_workspace_before_asynchronous_physical_cleanup(configured):
    other = jobs.create_workspace(configured["request"], "sample")
    sessions.revoke(configured["request"])
    session_id = hashlib.sha256(configured["nonce"].encode()).hexdigest()
    jobs.purge_session(session_id)
    for workspace_id in (
        configured["workspace"]["workspace_id"],
        other["workspace_id"],
    ):
        assert store.get("workspace", workspace_id)["state"] == "purged"
        assert ("cleanup", workspace_id, None) in configured["deliveries"]


def test_snapshot_manifest_bytes_are_counted_in_storage_reservation(
    configured, monkeypatch
):
    stream_stub(monkeypatch, configured, count=2)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    actual = sum(
        len(data) for generation, data in configured["objects"].values.values()
    )
    recorded = configured["control"].records["control:budget"]["reservations"][
        value["execution_id"]
    ]["bytes"]
    assert recorded == actual


def test_authenticated_download_can_coexist_with_same_user_page_read(
    configured, monkeypatch
):
    stream_stub(monkeypatch, configured, count=2)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    exported = jobs.create_export(
        configured["request"], "sample", value["workspace_id"], value["execution_id"]
    )
    jobs.run_export(exported["export_id"])
    output, size = jobs.open_download(
        configured["request"],
        "sample",
        value["workspace_id"],
        value["execution_id"],
        exported["export_id"],
    )
    with jobs._permit("read", "reader"):
        with pytest.raises(HTTPException):
            jobs.page(
                configured["request"],
                "sample",
                value["workspace_id"],
                value["execution_id"],
                0,
                100,
            )
    page = jobs.page(
        configured["request"],
        "sample",
        value["workspace_id"],
        value["execution_id"],
        0,
        100,
    )
    assert len(page["rows"]) == 2
    assert len(b"".join(output)) == size
    assert not configured["control"].records["control:budget"]["permits"]


def test_storage_admission_precedes_every_private_sql_upload(configured, monkeypatch):
    submit(configured)
    submit(configured, "SELECT 2", "second")
    stage = Mock(side_effect=results.stage_sql)
    monkeypatch.setattr(results, "stage_sql", stage)
    with pytest.raises(HTTPException) as denied:
        submit(configured, "SELECT 3", "third")
    assert denied.value.status_code == 429
    stage.assert_not_called()


def test_orphan_staging_cleanup_releases_pre_reserved_storage_only_after_delete(
    configured, monkeypatch
):
    identity = "d" * 32

    def reserve(values):
        control = jobs._coordinator(values["control:budget"])
        jobs._reserve(control, identity, "reader", results.SNAPSHOT_BYTES)
        return {"control:budget": control}

    store.mutate([("control", "budget")], reserve)
    results.stage_sql(identity, "SELECT 1")
    original = results.purge
    monkeypatch.setattr(
        results, "purge", Mock(side_effect=DatabaseUnavailable("outage"))
    )
    with pytest.raises(DatabaseUnavailable):
        jobs.cleanup(identity)
    assert identity in configured["control"].records["control:budget"]["reservations"]
    monkeypatch.setattr(results, "purge", original)
    jobs.cleanup(identity)
    assert (
        identity not in configured["control"].records["control:budget"]["reservations"]
    )


def test_uncertain_execution_creation_ack_preserves_durable_outbox_input(
    configured, monkeypatch
):
    backend = configured["control"]
    mutate = backend.mutate

    def uncertain(references, transform):
        mutate(references, transform)
        if any(kind == "request" for kind, identity in references):
            raise DatabaseUnavailable("ACK uncertain")

    monkeypatch.setattr(backend, "mutate", uncertain)
    with pytest.raises(DatabaseUnavailable):
        submit(configured)
    value = store.find("execution", "state", "queued", limit=1)[0]
    assert value["input"]["object"] in configured["objects"].values
    assert value["execution_id"] in backend.records["control:budget"]["reservations"]


def test_workspace_budget_is_atomic_even_for_concurrent_tabs(configured):
    outcomes = []

    def create():
        try:
            jobs.create_workspace(configured["request"], "sample")
            outcomes.append("created")
        except HTTPException as denied:
            outcomes.append(denied.status_code)

    threads = [threading.Thread(target=create) for _ in range(40)]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()
    assert outcomes.count("created") == 31 and outcomes.count(429) == 9
    session_id = hashlib.sha256(configured["nonce"].encode()).hexdigest()
    assert len(store.get("session", session_id)["workspace_ids"]) == 32


def test_partial_cleanup_deadline_keeps_reservation_and_retry_finishes_without_sql(
    configured, monkeypatch
):
    stream = stream_stub(monkeypatch, configured, count=1205)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    backend = configured["objects"]
    identity = value["execution_id"]
    configured["control"].records[f"execution:{identity}"]["expires_at"] = (
        time.time() - 1
    )
    clock = [0.0]
    monkeypatch.setattr(results.time, "monotonic", lambda: clock[0])
    original = backend.purge

    def partial(identity):
        first = next(
            path for path in backend.values if path.startswith(f"results/{identity}/")
        )
        backend.values.pop(first)
        clock[0] += 231

    monkeypatch.setattr(backend, "purge", partial)
    with pytest.raises(DatabaseUnavailable):
        jobs.cleanup(identity)
    assert backend.contains(identity)
    assert identity in configured["control"].records["control:budget"]["reservations"]
    assert not store.get("execution", identity).get("cleanup_done", False)
    monkeypatch.setattr(backend, "purge", original)
    jobs.cleanup(identity)
    assert not backend.contains(identity)
    assert (
        identity not in configured["control"].records["control:budget"]["reservations"]
    )
    assert store.get("execution", identity)["cleanup_done"]
    stream.assert_called_once()


def test_workspace_export_cas_cleanup_budget_stops_new_rpcs_and_preserves_quota_for_retry(
    configured, monkeypatch
):
    stream = stream_stub(monkeypatch, configured, count=2)
    value = submit(configured)
    jobs.run_query(value["execution_id"])
    identity, workspace_id = value["execution_id"], value["workspace_id"]
    budget = configured["control"].records["control:budget"]
    for index in range(8):
        export_id = f"{100 + index:032x}"
        artifact = results.put(
            f"results/{identity}/export-{export_id}.xlsx", b"PK-synthetic"
        )
        configured["control"].records[f"export:{export_id}"] = {
            "kind": "export",
            "export_id": export_id,
            "execution_id": identity,
            "workspace_id": workspace_id,
            "state": "completed",
            "worker_deadline": 0,
            "expires_at": time.time() + 3600,
            "file": artifact,
        }
        budget["reservations"][export_id] = {
            "owner": "reader",
            "bytes": artifact["bytes"],
        }
    clock, rpc_starts = [200.0], []
    monkeypatch.setattr(jobs.time, "monotonic", lambda: clock[0])
    original = configured["control"].mutate

    def slow_cas(references, transform):
        rpc_starts.append(clock[0])
        original(references, transform)
        clock[0] += 9

    monkeypatch.setattr(configured["control"], "mutate", slow_cas)
    purge = Mock(side_effect=results.purge)
    monkeypatch.setattr(results, "purge", purge)
    with pytest.raises(DatabaseUnavailable):
        jobs.purge_workspace_id(workspace_id, deadline=230)
    assert rpc_starts and all(start < 230 for start in rpc_starts)
    assert clock[0] == 236
    purge.assert_not_called()
    assert len(configured["control"].records["control:budget"]["reservations"]) == 9
    assert not store.get("workspace", workspace_id).get("cleanup_done", False)
    jobs.purge_workspace_id(workspace_id, deadline=clock[0] + 230)
    assert not configured["objects"].contains(identity)
    assert not configured["control"].records["control:budget"]["reservations"]
    assert store.get("workspace", workspace_id)["cleanup_done"]
    stream.assert_called_once()


def test_cleanup_cannot_start_metadata_rpc_after_budget_deadline(
    configured, monkeypatch
):
    monkeypatch.setattr(jobs.time, "monotonic", lambda: 230)
    get = Mock(side_effect=store.get)
    monkeypatch.setattr(store, "get", get)
    with pytest.raises(DatabaseUnavailable):
        jobs.cleanup("d" * 32, deadline=230)
    get.assert_not_called()
