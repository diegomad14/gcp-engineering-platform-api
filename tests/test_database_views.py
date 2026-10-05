"""Global order consumes one capture, including paging/Excel and revocation."""

import io
from unittest.mock import Mock

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest

from eng_platform_api.services import database_console as console
from eng_platform_api.services import database_jobs as jobs
from eng_platform_api.services import database_sessions as sessions
from eng_platform_api.services import database_views as views
from eng_platform_api.services.database_registry import DatabaseUnavailable
from tests.test_database_jobs import configured as configured, HEADERS, browser, submit


def captured(environment, monkeypatch, rows=None):
    rows = (
        rows
        if rows is not None
        else [[index, str(index) + ".1234567890123456789"] for index in range(1205)]
    )
    columns = [
        {"key": "0", "name": "duplicate", "data_type": "int8"},
        {"key": "1", "name": "duplicate", "data_type": "numeric"},
    ]

    def stream(
        database, statement, authorize, on_columns, on_rows, *, admission, **kwargs
    ):
        with admission:
            authorize()
            on_columns(columns)
            for start in range(0, len(rows), 16):
                on_rows(rows[start : start + 16])
        return {"row_count": len(rows), "elapsed_ms": 1}

    call = Mock(side_effect=stream)
    monkeypatch.setattr(console, "stream_query", call)
    value = submit(environment)
    jobs.run_query(value["execution_id"])
    return value["execution_id"], call


def args(environment, execution_id):
    return (
        environment["request"],
        "sample",
        environment["workspace"]["workspace_id"],
        execution_id,
    )


def sorted_view(environment, execution_id, column="0", direction="desc", key="sort-1"):
    value = views.create_view(*args(environment, execution_id), column, direction, key)
    views.run_sort(value["view_id"])
    return views.view(*args(environment, execution_id), value["view_id"])


def test_full_sorted_snapshot_pages_and_excel_use_capture_once(configured, monkeypatch):
    identity, stream = captured(configured, monkeypatch)
    original = jobs.page(*args(configured, identity), 0, 100)
    selected = sorted_view(configured, identity)
    assert selected["status"] == "completed" and selected["row_count"] == 1205
    assert (
        selected["expires_at"]
        == jobs.execution(*args(configured, identity))["expires_at"]
    )
    all_rows = []
    for index in range(13):
        page = jobs.page(*args(configured, identity), index, 100, selected["view_id"])
        all_rows.extend(page["rows"])
        assert page["view_id"] == selected["view_id"] and page["total_pages"] == 13
    assert [row[0] for row in all_rows] == list(range(1204, -1, -1))
    assert jobs.page(*args(configured, identity), 0, 100) == original
    exported = jobs.create_export(*args(configured, identity), selected["view_id"])
    jobs.run_export(exported["export_id"])
    metadata = jobs.export(*args(configured, identity), exported["export_id"])
    assert metadata["source_view_id"] == selected["view_id"]
    file, size = jobs.open_download(*args(configured, identity), exported["export_id"])
    raw = b"".join(file)
    assert len(raw) == size
    workbook = load_workbook(io.BytesIO(raw), read_only=True)
    excel = list(workbook.active.values)
    assert len(excel) == 1206 and excel[1] == (1204, "1204.1234567890123456789")
    assert excel[-1] == (0, "0.1234567890123456789")
    stream.assert_called_once()
    assert not any("sort-" in path for path in configured["objects"].values)


@pytest.mark.parametrize(
    "direction,order", [("asc", [2, 1, 0, 3]), ("desc", [0, 1, 2, 3])]
)
def test_exact_decimals_and_null_last_in_both_directions(
    configured, monkeypatch, direction, order
):
    rows = [
        [0, "9007199254740993.0000000000000000002"],
        [1, "9007199254740993.0000000000000000001"],
        [2, "-1e-40"],
        [3, None],
    ]
    identity, _ = captured(configured, monkeypatch, rows)
    selected = sorted_view(configured, identity, "1", direction)
    page = jobs.page(*args(configured, identity), 0, 25, selected["view_id"])
    assert [row[0] for row in page["rows"]] == order


def test_idempotency_conflict_and_duplicate_task_never_rerun_capture(
    configured, monkeypatch
):
    identity, stream = captured(configured, monkeypatch)
    first = views.create_view(*args(configured, identity), "0", "asc", "same")
    assert (
        views.create_view(*args(configured, identity), "0", "asc", "same")["view_id"]
        == first["view_id"]
    )
    with pytest.raises(HTTPException) as error:
        views.create_view(*args(configured, identity), "0", "desc", "same")
    assert error.value.status_code == 409
    views.run_sort(first["view_id"])
    objects = dict(configured["objects"].values)
    views.run_sort(first["view_id"])
    assert configured["objects"].values == objects
    stream.assert_called_once()


def test_quota_failure_preserves_prior_view_and_capture(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    prior = sorted_view(configured, identity)
    previous = jobs.page(*args(configured, identity), 0, 25, prior["view_id"])
    used = sum(
        item["bytes"]
        for item in configured["control"]
        .get("control", "budget")["reservations"]
        .values()
    )
    monkeypatch.setattr(jobs, "USER_BYTES", used + 1)
    failed = views.create_view(*args(configured, identity), "0", "asc", "new")
    views.run_sort(failed["view_id"])
    record = configured["control"].get("view", failed["view_id"])
    assert record["state"] == "failed" and record["error_code"] == "RESOURCE_EXCEEDED"
    assert record["cleanup_done"] and not configured["objects"].contains(
        failed["view_id"]
    )
    assert jobs.page(*args(configured, identity), 0, 25, prior["view_id"]) == previous
    assert jobs.page(*args(configured, identity), 0, 25)["rows"][0][0] == 0


def test_cancel_queued_and_cleanup_preserve_original(configured, monkeypatch):
    identity, stream = captured(configured, monkeypatch)
    value = views.create_view(*args(configured, identity), "0", "desc", "cancel")
    assert views.delete_view(*args(configured, identity), value["view_id"]) == {
        "purged": True
    }
    views.run_sort(value["view_id"])
    jobs.cleanup(value["view_id"])
    assert (
        value["view_id"]
        not in configured["control"].get("execution", identity)["view_ids"]
    )
    assert (
        value["view_id"]
        not in configured["control"].get("control", "budget")["reservations"]
    )
    assert jobs.page(*args(configured, identity), 0, 25)["row_count"] == 1205
    stream.assert_called_once()


def test_live_view_budget_recovers_after_deletion(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch, [[1, "1"]])
    monkeypatch.setattr(views, "MAX_VIEWS", 2)
    for index in range(20):
        value = views.create_view(*args(configured, identity), "0", "asc", str(index))
        views.delete_view(*args(configured, identity), value["view_id"])
        jobs.cleanup(value["view_id"])
    assert configured["control"].get("execution", identity)["view_ids"] == []


def test_logout_prevents_publish_and_purges_view_children(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    value = sorted_view(configured, identity)
    export = jobs.create_export(*args(configured, identity), value["view_id"])
    jobs.run_export(export["export_id"])
    sessions.revoke(configured["request"])
    jobs.purge_workspace_id(configured["workspace"]["workspace_id"])
    assert not configured["objects"].contains(identity)
    assert not configured["objects"].contains(value["view_id"])
    assert configured["control"].get("control", "budget")["reservations"] == {}
    with pytest.raises(HTTPException):
        jobs.page(*args(configured, identity), 0, 25, value["view_id"])


def test_expiry_and_permission_changes_fail_closed(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    value = sorted_view(configured, identity)
    source = configured["control"].records[f"execution:{identity}"]
    source["expires_at"] = 0
    with pytest.raises(HTTPException) as error:
        jobs.page(*args(configured, identity), 0, 25, value["view_id"])
    assert error.value.status_code == 410
    jobs.cleanup(identity)
    assert not configured["objects"].contains(value["view_id"])


def test_storage_credit_retained_until_namespace_verified_empty(
    configured, monkeypatch
):
    identity, _ = captured(configured, monkeypatch)
    value = sorted_view(configured, identity)
    views.delete_view(*args(configured, identity), value["view_id"])
    original = configured["objects"].purge
    monkeypatch.setattr(configured["objects"], "purge", lambda identity: None)
    with pytest.raises(DatabaseUnavailable):
        jobs.cleanup(value["view_id"])
    assert (
        value["view_id"]
        in configured["control"].get("control", "budget")["reservations"]
    )
    monkeypatch.setattr(configured["objects"], "purge", original)
    jobs.cleanup(value["view_id"])
    assert (
        value["view_id"]
        not in configured["control"].get("control", "budget")["reservations"]
    )


def test_routes_202_view_selection_capability_and_private_guards(
    configured, monkeypatch
):
    identity, _ = captured(configured, monkeypatch)
    prefix = f"/api/databases/sample/workspaces/{configured['workspace']['workspace_id']}/executions/{identity}"
    client = browser(configured)
    assert client.get("/api/databases").json()["global_sort_enabled"]
    payload = {"column_key": "0", "direction": "desc", "client_request_id": "routes"}
    assert client.post(prefix + "/views", json=payload).status_code == 403
    assert (
        client.post(
            prefix + "/views", json={**payload, "extra": True}, headers=HEADERS
        ).status_code
        == 422
    )
    created = client.post(prefix + "/views", json=payload, headers=HEADERS)
    assert created.status_code == 202 and created.headers["cache-control"] == "no-store"
    view_id = created.json()["view_id"]
    views.run_sort(view_id)
    assert client.get(prefix + f"/views/{view_id}").json()["status"] == "completed"
    page = client.post(
        prefix + "/pages",
        json={"view_id": view_id, "page_index": 0, "page_size": 25},
        headers=HEADERS,
    )
    assert page.status_code == 200 and page.json()["rows"][0][0] == 1204
    exported = client.post(
        prefix + "/exports", json={"view_id": view_id}, headers=HEADERS
    )
    assert exported.status_code == 202 and exported.json()["source_view_id"] == view_id
    assert (
        client.request(
            "DELETE", prefix + f"/views/{view_id}", headers=HEADERS, json={}
        ).status_code
        == 200
    )


def test_completed_excel_survives_retiring_its_view(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    selected = sorted_view(configured, identity)
    exported = jobs.create_export(*args(configured, identity), selected["view_id"])
    jobs.run_export(exported["export_id"])
    views.delete_view(*args(configured, identity), selected["view_id"])
    jobs.cleanup(selected["view_id"])
    assert not configured["objects"].contains(selected["view_id"])
    stream, _ = jobs.open_download(*args(configured, identity), exported["export_id"])
    workbook = load_workbook(io.BytesIO(b"".join(stream)), read_only=True)
    assert next(iter(workbook.active.values)) == ("duplicate", "duplicate")
    assert (
        jobs.export(*args(configured, identity), exported["export_id"])["status"]
        == "completed"
    )


def test_revoke_at_publish_fence_rejects_late_complete_view(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    selected = views.create_view(*args(configured, identity), "0", "desc", "fence")
    finish = jobs._finish

    def revoke_before_publish(identity, claim, state, **kwargs):
        if state == "completed" and kwargs.get("kind") == "view":
            sessions.revoke(configured["request"])
        return finish(identity, claim, state, **kwargs)

    monkeypatch.setattr(jobs, "_finish", revoke_before_publish)
    views.run_sort(selected["view_id"])
    record = configured["control"].get("view", selected["view_id"])
    assert record["state"] == "cancelled" and record["manifest"] is None
    assert record["cleanup_done"] and not configured["objects"].contains(
        selected["view_id"]
    )


def test_cancel_running_after_run_upload_prevents_publication(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    selected = views.create_view(*args(configured, identity), "0", "desc", "during")
    write = views.SortObjects.write_run

    def upload_then_cancel(adapter, name, pieces):
        result = write(adapter, name, pieces)
        views.delete_view(*args(configured, identity), selected["view_id"])
        return result

    monkeypatch.setattr(views.SortObjects, "write_run", upload_then_cancel)
    views.run_sort(selected["view_id"])
    assert configured["control"].get("view", selected["view_id"])["manifest"] is None
    assert not configured["objects"].contains(selected["view_id"])
    assert jobs.page(*args(configured, identity), 0, 25)["row_count"] == 1205


def test_sort_and_export_share_one_global_heavy_permit(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    selected = views.create_view(*args(configured, identity), "0", "asc", "busy")
    with jobs._permit("export", "other"):
        views.run_sort(selected["view_id"])
    value = views.view(*args(configured, identity), selected["view_id"])
    assert value["status"] == "failed" and value["error_code"] == "RESOURCE_EXCEEDED"


def test_retiring_running_export_source_stops_export_only(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    selected = sorted_view(configured, identity)
    exported = jobs.create_export(*args(configured, identity), selected["view_id"])
    views.delete_view(*args(configured, identity), selected["view_id"])
    jobs.run_export(exported["export_id"])
    jobs.cleanup(selected["view_id"])
    assert configured["control"].get("export", exported["export_id"])["state"] in {
        "failed",
        "purged",
        "expired",
    }
    assert jobs.page(*args(configured, identity), 0, 25)["rows"][0][0] == 0


def test_queued_outbox_timeout_and_lost_worker_cleanup_are_visible(
    configured, monkeypatch
):
    identity, _ = captured(configured, monkeypatch)
    queued = views.create_view(*args(configured, identity), "0", "asc", "outbox")
    configured["control"].records[f"view:{queued['view_id']}"]["dispatch_deadline"] = 0
    assert (
        views.view(*args(configured, identity), queued["view_id"])["error_code"]
        == "DISPATCH_FAILED"
    )
    running = views.create_view(*args(configured, identity), "0", "desc", "lost")
    jobs._claim("view", running["view_id"])
    configured["control"].records[f"view:{running['view_id']}"]["worker_deadline"] = 0
    views.sweep_views(jobs.time.monotonic() + 230)
    assert (
        views.view(*args(configured, identity), running["view_id"])["error_code"]
        == "WORKER_LOST"
    )


@pytest.mark.parametrize("column,direction", [("unknown", "asc"), ("0", "invalid")])
def test_invalid_sort_parameters_never_enqueue_or_upload(
    configured, monkeypatch, column, direction
):
    identity, _ = captured(configured, monkeypatch)
    before = dict(configured["objects"].values)
    with pytest.raises(HTTPException) as error:
        views.create_view(*args(configured, identity), column, direction, "invalid")
    assert error.value.status_code == 422 and configured["objects"].values == before


def test_sort_run_adapter_generation_integrity_and_accounting(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch, [[1, "1"]])
    selected = views.create_view(*args(configured, identity), "0", "asc", "adapter")
    claim = jobs._claim("view", selected["view_id"])
    adapter = views.SortObjects(jobs.RunContext(claim, kind="view"))
    raw = b'{"ordinal":0,"row":[1,"1"]}\n'
    run = adapter.write_run("run-000000.ndjson", iter([raw]))
    assert list(adapter.read_run(run)) == [raw]
    assert configured["control"].get("control", "budget")["reservations"][
        selected["view_id"]
    ]["bytes"] == len(raw)
    adapter.delete_run(run)
    assert (
        configured["control"].get("control", "budget")["reservations"][
            selected["view_id"]
        ]["bytes"]
        == 0
    )
    with pytest.raises(DatabaseUnavailable):
        adapter.resize(-1)
    for name, pieces in [
        ("../invalid", iter([raw])),
        ("run-000001.ndjson", iter([b"x"])),
    ]:
        with pytest.raises(DatabaseUnavailable):
            adapter.write_run(name, pieces)


def test_cleanup_deadline_does_not_credit_partial_storage(configured, monkeypatch):
    identity, _ = captured(configured, monkeypatch)
    selected = sorted_view(configured, identity)
    views.delete_view(*args(configured, identity), selected["view_id"])
    with pytest.raises(DatabaseUnavailable):
        views.cleanup_view(selected["view_id"], deadline=0)
    assert configured["objects"].contains(selected["view_id"])
    assert (
        selected["view_id"]
        in configured["control"].get("control", "budget")["reservations"]
    )
    jobs.cleanup(selected["view_id"])


def test_incomplete_capture_and_foreign_view_are_rejected(configured, monkeypatch):
    pending = submit(configured)
    with pytest.raises(HTTPException) as error:
        views.create_view(
            *args(configured, pending["execution_id"]), "0", "asc", "not-ready"
        )
    assert error.value.status_code == 409
    identity, _ = captured(configured, monkeypatch, [[1, "1"]])
    with pytest.raises(HTTPException) as missing:
        views.require_view(*args(configured, identity), "f" * 32)
    assert missing.value.status_code == 404
