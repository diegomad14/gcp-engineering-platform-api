"""Read exported workbooks independently; exercise streaming safety budgets."""

from io import BytesIO
import json
from unittest.mock import Mock
from xml.etree import ElementTree
import zipfile

import openpyxl
import pytest

from eng_platform_api.services import database_xlsx as xlsx
from eng_platform_api.services.database_console import DatabaseResourceExceeded


def exported(columns, rows, **kwargs):
    output = BytesIO()
    result = xlsx.write_xlsx(output, columns, rows, **kwargs)
    output.seek(0)
    return (
        openpyxl.load_workbook(output, read_only=True, data_only=False),
        output,
        result,
    )


def column(name, data_type="text"):
    return {"name": name, "data_type": data_type}


def test_independent_reader_preserves_exact_values_literal_text_and_duplicates():
    headers = [
        column("duplicate", "int4"),
        column("duplicate", "int8"),
        column("amount", "numeric"),
        column("formula"),
        column("unicode"),
        column("payload", "jsonb"),
        column("bool", "bool"),
        column("date", "timestamptz"),
    ]
    data = [
        [
            1,
            "9007199254740993",
            "1234567890.12345678901234567890",
            '=HYPERLINK("https://evil.invalid")',
            "Árbol 🛰️ <&>",
            {"big": "9007199254740993", "list": [1, True]},
            True,
            "2026-10-04T21:33:00-05:00",
        ]
    ]
    book, output, result = exported(headers, data)
    sheet = book.worksheets[0]
    rows = list(sheet.iter_rows())
    assert [cell.value for cell in rows[0]] == [item["name"] for item in headers]
    assert [cell.value for cell in rows[1]] == [
        1,
        data[0][1],
        data[0][2],
        data[0][3],
        data[0][4],
        json.dumps(data[0][5], ensure_ascii=False, separators=(",", ":")),
        True,
        data[0][7],
    ]
    assert rows[1][3].data_type == "s"
    assert result["row_count"] == 1 and result["sheet_count"] == 1
    assert result["bytes"] == len(output.getvalue())
    with zipfile.ZipFile(output) as archive:
        assert not any(
            "sharedStrings" in name or "vba" in name for name in archive.namelist()
        )
        xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"<f>" not in xml and b"hyperlink" not in xml
        for name in archive.namelist():
            ElementTree.fromstring(archive.read(name))
    book.close()


def test_null_empty_and_literal_null_marker_are_reversible():
    book, _, _ = exported(
        [column("value")],
        [
            [None],
            [""],
            ["\\N"],
            ["\\path"],
            ["NULL"],
            ["+123"],
            ["@name"],
            ["-not-a-number"],
        ],
    )
    values = [row[0].value for row in list(book.active.iter_rows())[1:]]
    assert values == [
        "\\N",
        "",
        "\\\\N",
        "\\\\path",
        "NULL",
        "+123",
        "@name",
        "-not-a-number",
    ]
    book.close()


def test_canonical_json_preserves_nested_numeric_types_and_json_null():
    document = (
        '{"amount":0.12345678901234567890,"id":9007199254740993,"unicode":"Árbol 🛰"}'
    )
    book, _, _ = exported(
        [column("payload", "jsonb"), column("array", "_numeric")],
        [[document, "[0.12345678901234567890,9007199254740993]"], ["null", None]],
    )
    rows = list(book.active.values)
    assert rows[1] == (document, "[0.12345678901234567890,9007199254740993]")
    assert rows[2] == ("null", "\\N")
    assert isinstance(json.loads(rows[1][0])["id"], int)
    book.close()


def test_empty_result_keeps_header_and_numeric_precision_is_not_inferred():
    book, _, result = exported([column("value", "int8")], [])
    assert list(book.active.values) == [("value",)]
    assert result["row_count"] == 0
    book.close()
    book, _, _ = exported(
        [
            column("value", "int8"),
            column("float", "float8"),
            column("numeric", "numeric"),
        ],
        [[999999999999999, "1.2345678901234567", 1], [1000000000000000, "Infinity", 2]],
    )
    rows = list(book.active.values)
    assert rows[1] == (999999999999999, "1.2345678901234567", "1")
    assert rows[2] == ("1000000000000000", "Infinity", "2")
    book.close()


def test_sheet_split_preserves_every_row_and_empty_final_sheet_is_avoided(monkeypatch):
    monkeypatch.setattr(xlsx, "MAX_DATA_ROWS", 3)
    for count, sheets in [(0, 1), (3, 1), (4, 2), (6, 2), (7, 3)]:
        book, _, result = exported(
            [column("id", "int4")], ([index] for index in range(count))
        )
        values = [row[0] for sheet in book for row in list(sheet.values)[1:]]
        assert values == list(range(count))
        assert result["row_count"] == count and result["sheet_count"] == sheets
        book.close()


def test_xml_controls_and_literal_office_escape_sequences_are_lossless_encoded():
    _, output, _ = exported([column("controls")], [["a\x01\rb_x0001_\t\n"]])
    with zipfile.ZipFile(output) as archive:
        xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"a_x0001__x000D_b_x005F_x0001_\t\n" in xml
        assert b"\x01" not in xml
        ElementTree.fromstring(xml)


@pytest.mark.parametrize(
    "value",
    ["x" * 32768, "🛰" * 16384, "\n" * 254, "\ud800", float("nan"), float("inf")],
)
def test_incompatible_cell_fails_instead_of_truncating(value):
    with pytest.raises(DatabaseResourceExceeded):
        xlsx.write_xlsx(BytesIO(), [column("value")], [[value]])


def test_columns_and_row_shape_fail_closed():
    for columns in [[], [column("x")] * 101, [{"name": 1, "data_type": "int4"}]]:
        with pytest.raises(DatabaseResourceExceeded):
            xlsx.write_xlsx(BytesIO(), columns, [])
    with pytest.raises(DatabaseResourceExceeded):
        xlsx.write_xlsx(BytesIO(), [column("x")], [[1, 2]])


def test_permission_is_rechecked_before_finalization_and_no_error_contains_values():
    calls = Mock(side_effect=[None, PermissionError("revoked")])
    with pytest.raises(PermissionError):
        xlsx.write_xlsx(BytesIO(), [column("x")], [["PRIVATE"]], calls)
    assert calls.call_count == 2


def test_time_and_compressed_xml_budgets_abort_stream(monkeypatch):
    with pytest.raises(DatabaseResourceExceeded):
        xlsx.write_xlsx(BytesIO(), [column("x")], [], timeout_seconds=241)
    times = iter([0, 0])
    monkeypatch.setattr(xlsx, "monotonic", lambda: next(times, 241))
    with pytest.raises(DatabaseResourceExceeded):
        xlsx.write_xlsx(BytesIO(), [column("x")], [[1]])


def test_output_caps_and_short_writes(monkeypatch):
    monkeypatch.setattr(xlsx, "MAX_COMPRESSED_BYTES", 10)
    with pytest.raises(DatabaseResourceExceeded):
        xlsx.write_xlsx(BytesIO(), [column("x")], [[1]])
    monkeypatch.setattr(xlsx, "MAX_COMPRESSED_BYTES", 1000000)
    monkeypatch.setattr(xlsx, "MAX_XML_BYTES", 10)
    with pytest.raises(DatabaseResourceExceeded):
        xlsx.write_xlsx(BytesIO(), [column("x")], [[1]])
    guard = xlsx._Guard(lambda: None, 240)
    sink = Mock()
    sink.write.return_value = 0
    with pytest.raises(OSError):
        xlsx._Output(sink, guard).write(b"x")


def test_stream_uses_one_pass_input_and_bounded_memory(tmp_path):
    import tracemalloc

    def source():
        for index in range(10000):
            yield [index, "text" * 300]

    tracemalloc.start()
    try:
        with (tmp_path / "stream.xlsx").open("wb") as sink:
            result = xlsx.write_xlsx(
                sink, [column("id", "int4"), column("text")], source()
            )
        _, peak = tracemalloc.get_traced_memory()
    finally:
        tracemalloc.stop()
    assert result["row_count"] == 10000
    assert peak < 8 * 1024 * 1024
    book = openpyxl.load_workbook(tmp_path / "stream.xlsx", read_only=True)
    assert sum(1 for _ in book.active.values) == 10001
    book.close()
